#!/usr/bin/env python3
"""Servier-style PPTX: positive/negative cellular interactions regulating IL-17.

Shows interacting cells (Th17, stroma, DC, Treg, M2, neutrophils…) with
molecular contact points (podoplanin, truncated CD74, CLEC2).
Based on Noack & Miossec.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

OUT = "/workspace/figures/il17/Figure_IL17_interactions_cellulaires.pptx"

BG = RGBColor(0xF8, 0xFA, 0xFC)
INK = RGBColor(0x2C, 0x3E, 0x50)
MUTED = RGBColor(0x5D, 0x6D, 0x7E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xCB, 0xD5, 0xE1)

POS = RGBColor(0x0D, 0x94, 0x88)
POS_DK = RGBColor(0x0A, 0x6F, 0x66)
POS_LT = RGBColor(0xE0, 0xF5, 0xF2)
POS_BAND = RGBColor(0xE8, 0xF7, 0xF4)

NEG = RGBColor(0xD4, 0x6A, 0x4E)
NEG_DK = RGBColor(0xA8, 0x4A, 0x32)
NEG_LT = RGBColor(0xFB, 0xEE, 0xE9)
NEG_BAND = RGBColor(0xFC, 0xF1, 0xEC)

CORE = RGBColor(0xC4, 0x45, 0x69)
CORE_DK = RGBColor(0x8E, 0x2D, 0x4A)
CORE_LT = RGBColor(0xFB, 0xE9, 0xEF)

# Cell type colors (Servier-like soft fills)
TH17 = RGBColor(0xE8, 0x5D, 0x75)
TH17_LT = RGBColor(0xFB, 0xD5, 0xDC)
STROMA = RGBColor(0x3A, 0x7C, 0xA5)
STROMA_LT = RGBColor(0xD6, 0xE8, 0xF4)
DC = RGBColor(0x8B, 0x5C, 0xF6)
DC_LT = RGBColor(0xE9, 0xDE, 0xFF)
MAC = RGBColor(0xD9, 0x77, 0x06)
MAC_LT = RGBColor(0xFE, 0xE9, 0xC8)
GDT = RGBColor(0x0E, 0xA5, 0xE9)
GDT_LT = RGBColor(0xD0, 0xEF, 0xFC)
ILC = RGBColor(0x14, 0xB8, 0xA6)
ILC_LT = RGBColor(0xCC, 0xF5, 0xEF)
TREG = RGBColor(0x65, 0xA3, 0x0D)
TREG_LT = RGBColor(0xE2, 0xF2, 0xC8)
M2 = RGBColor(0x4D, 0x7C, 0x0F)
M2_LT = RGBColor(0xDC, 0xED, 0xC8)
NEU = RGBColor(0x64, 0x74, 0x8B)
NEU_LT = RGBColor(0xE2, 0xE8, 0xF0)
MOL = RGBColor(0x7C, 0x3A, 0xED)
MOL_LT = RGBColor(0xF1, 0xEA, 0xFE)


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def set_line(shape, color, width_pt=1.15):
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def round_rect(slide, left, top, width, height, fill, line=None, adj=0.12):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    set_fill(shape, fill)
    if line is None:
        no_line(shape)
    else:
        set_line(shape, line)
    try:
        shape.adjustments[0] = adj
    except Exception:
        pass
    return shape


def oval(slide, left, top, width, height, fill, line=None, line_w=1.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    set_fill(shape, fill)
    if line is None:
        no_line(shape)
    else:
        set_line(shape, line, line_w)
    return shape


def arrow_right(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    set_fill(shape, color)
    no_line(shape)
    return shape


def arrow_left(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, left, top, width, height)
    set_fill(shape, color)
    no_line(shape)
    return shape


def arrow_down(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    set_fill(shape, color)
    no_line(shape)
    return shape


def set_runs(shape, lines, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    try:
        tf._txBody.get_or_add_bodyPr().set("anchor", "ctr")
    except Exception:
        pass
    for i, line in enumerate(lines):
        if isinstance(line, str):
            text, size, bold, color = line, 11, False, INK
        else:
            text, size, bold, color = line
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return tf


def textbox(slide, left, top, width, height, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def cell(slide, left, top, w, h, fill, border, name, sub=None, name_size=11, sub_size=8):
    """Draw a cell as an oval with label (Servier-like)."""
    shape = oval(slide, left, top, w, h, fill, border, line_w=1.75)
    lines = [(name, name_size, True, border)]
    if sub:
        lines.append((sub, sub_size, False, MUTED))
    set_runs(shape, lines)
    return shape


def contact_bridge(slide, left, top, width, height, label, fill=MOL_LT, accent=MOL):
    """Small molecular contact label between cells."""
    shape = round_rect(slide, left, top, width, height, fill, accent, adj=0.25)
    set_runs(shape, [(label, 8, True, accent)])
    return shape


def molecule_tag(slide, left, top, width, height, text, fill, accent):
    shape = round_rect(slide, left, top, width, height, fill, accent, adj=0.2)
    set_runs(shape, [(text, 8.5, True, accent)])
    return shape


def build_schema_slide(prs, lang="fr"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    set_fill(bg, BG)
    no_line(bg)

    if lang == "fr":
        title = "Régulation de l’IL-17 par interactions cellulaires"
        subtitle = "Régulation positive (gauche) et négative (droite) — cellules en contact"
        pos_hdr = "+  INTERACTIONS ACTIVATRICES"
        neg_hdr = "−  INTERACTIONS INHIBITRICES"
        center_title = "Production d’IL-17"
        mol_title = "Molécules de contact"
        footer = (
            "D’après Noack & Miossec — Interactions cellulaires et moléculaires "
            "(podoplanine, CD74 tronqué, CLEC2, Treg/M2)  |  Style Servier"
        )
        stroma_name, stroma_sub = "Synoviocyte", "fibroblaste / stroma"
        th17_sub = "RORγt⁺"
        dc_sub = "IL-1β · IL-6 · IL-23"
        mac_sub = "inflammation"
        gdt_sub = "source précoce"
        ilc_sub = "barrières"
        treg_sub = "FoxP3⁺"
        m2_sub = "IL-10 · TGF-β"
        neu_sub = "apoptose → ↓IL-23"
        note_pos = "Contact stroma–immunité ↑ IL-17"
        note_neg = "Treg / M2 / feedback ↓ IL-17"
        out_label = "IL-17A / F"
        clec_note = "pdpn–CLEC2 : ↓IL-17 ↑IL-10"
    else:
        title = "Cellular interactions regulating IL-17"
        subtitle = "Positive (left) and negative (right) regulation — interacting cells"
        pos_hdr = "+  ACTIVATING INTERACTIONS"
        neg_hdr = "−  INHIBITORY INTERACTIONS"
        center_title = "IL-17 production"
        mol_title = "Contact molecules"
        footer = (
            "Based on Noack & Miossec — Cellular and molecular interactions "
            "(podoplanin, truncated CD74, CLEC2, Treg/M2)  |  Servier style"
        )
        stroma_name, stroma_sub = "Synoviocyte", "fibroblast / stroma"
        th17_sub = "RORγt⁺"
        dc_sub = "IL-1β · IL-6 · IL-23"
        mac_sub = "inflammation"
        gdt_sub = "early source"
        ilc_sub = "barrier tissues"
        treg_sub = "FoxP3⁺"
        m2_sub = "IL-10 · TGF-β"
        neu_sub = "apoptosis → ↓IL-23"
        note_pos = "Stroma–immune contact ↑ IL-17"
        note_neg = "Treg / M2 / feedback ↓ IL-17"
        out_label = "IL-17A / F"
        clec_note = "pdpn–CLEC2: ↓IL-17 ↑IL-10"

    textbox(
        slide, Inches(0.25), Inches(0.06), Inches(12.8), Inches(0.32),
        title, size=19, bold=True, color=INK, align=PP_ALIGN.CENTER,
    )
    textbox(
        slide, Inches(0.25), Inches(0.34), Inches(12.8), Inches(0.22),
        subtitle, size=11, color=MUTED, align=PP_ALIGN.CENTER,
    )

    # ===== LEFT: positive interactions =====
    round_rect(slide, Inches(0.18), Inches(0.62), Inches(5.35), Inches(6.45), POS_BAND, POS, adj=0.03)
    pos_title = round_rect(slide, Inches(0.55), Inches(0.72), Inches(4.6), Inches(0.30), POS, None, adj=0.2)
    set_runs(pos_title, [(pos_hdr, 11, True, WHITE)])

    # Interaction pair 1: Th17 / γδ / ILC3  ↔  synoviocyte
    pair1 = round_rect(slide, Inches(0.35), Inches(1.12), Inches(5.00), Inches(2.35), WHITE, POS, adj=0.06)
    textbox(slide, Inches(0.45), Inches(1.18), Inches(4.8), Inches(0.22), note_pos, size=9, bold=True, color=POS_DK, align=PP_ALIGN.CENTER)

    # Immune cluster
    cell(slide, Inches(0.55), Inches(1.50), Inches(1.15), Inches(0.85), TH17_LT, TH17, "Th17", th17_sub, 12, 8)
    cell(slide, Inches(0.55), Inches(2.45), Inches(0.95), Inches(0.70), GDT_LT, GDT, "γδ T", gdt_sub, 11, 7.5)
    cell(slide, Inches(1.60), Inches(2.45), Inches(0.95), Inches(0.70), ILC_LT, ILC, "ILC3", ilc_sub, 11, 7.5)

    # Contact zone with molecules
    contact_bridge(slide, Inches(2.65), Inches(1.70), Inches(1.10), Inches(0.32), "podoplanine")
    contact_bridge(slide, Inches(2.65), Inches(2.10), Inches(1.10), Inches(0.32), "CD74 trunc.")
    # dashed-like contact dots
    for dy in (0.0, 0.18, 0.36):
        d = oval(slide, Inches(2.55), Inches(1.78) + Inches(dy), Inches(0.08), Inches(0.08), POS, None)
        d2 = oval(slide, Inches(3.75), Inches(1.78) + Inches(dy), Inches(0.08), Inches(0.08), POS, None)

    # Stromal cell (larger elongated = fibroblast-like)
    stroma = round_rect(slide, Inches(3.95), Inches(1.55), Inches(1.25), Inches(1.55), STROMA_LT, STROMA, adj=0.35)
    set_runs(stroma, [(stroma_name, 11, True, STROMA), (stroma_sub, 7.5, False, MUTED)])

    molecule_tag(slide, Inches(3.95), Inches(3.18), Inches(1.25), Inches(0.22), "caspase-1", POS_LT, POS)

    # Interaction pair 2: DC / macrophage feeding Th17
    pair2 = round_rect(slide, Inches(0.35), Inches(3.60), Inches(5.00), Inches(2.00), WHITE, POS, adj=0.06)
    textbox(
        slide, Inches(0.45), Inches(3.66), Inches(4.8), Inches(0.20),
        "DC / macrophages → cytokines Th17" if lang == "fr" else "DC / macrophages → Th17 cytokines",
        size=9, bold=True, color=POS_DK, align=PP_ALIGN.CENTER,
    )
    cell(slide, Inches(0.60), Inches(4.00), Inches(1.30), Inches(1.05), DC_LT, DC, "DC" if lang == "fr" else "DC", dc_sub, 12, 7.5)
    cell(slide, Inches(2.15), Inches(4.10), Inches(1.20), Inches(0.95), MAC_LT, MAC, "Macrophage", mac_sub, 10, 8)
    arrow_right(slide, Inches(3.45), Inches(4.40), Inches(0.40), Inches(0.22), POS)
    cell(slide, Inches(3.95), Inches(4.05), Inches(1.20), Inches(1.05), TH17_LT, TH17, "Th17", "↑ expansion", 12, 8)

    # Bottom positive note
    pos_note = round_rect(slide, Inches(0.35), Inches(5.75), Inches(5.00), Inches(1.15), POS_LT, POS, adj=0.08)
    if lang == "fr":
        set_runs(pos_note, [
            ("Boucle locale", 10, True, POS_DK),
            ("IL-17 → stroma : IL-6, CXCL1/8, CCL20", 9, False, INK),
            ("→ recrutement Th17 / neutrophiles", 9, False, INK),
            ("Endothélium : ICAM-1, VCAM-1", 9, False, INK),
        ])
    else:
        set_runs(pos_note, [
            ("Local loop", 10, True, POS_DK),
            ("IL-17 → stroma: IL-6, CXCL1/8, CCL20", 9, False, INK),
            ("→ Th17 / neutrophil recruitment", 9, False, INK),
            ("Endothelium: ICAM-1, VCAM-1", 9, False, INK),
        ])

    # ===== CENTER hub =====
    round_rect(slide, Inches(5.70), Inches(0.62), Inches(1.90), Inches(6.45), CORE_LT, CORE, adj=0.04)
    hub = round_rect(slide, Inches(5.85), Inches(0.80), Inches(1.60), Inches(0.55), CORE, None, adj=0.15)
    set_runs(hub, [(center_title, 10, True, WHITE)])

    # Central Th17 producing IL-17
    cell(slide, Inches(5.90), Inches(1.55), Inches(1.50), Inches(1.20), TH17_LT, TH17, "Th17", th17_sub, 13, 9)
    arrow_down(slide, Inches(6.45), Inches(2.85), Inches(0.28), Inches(0.35), CORE)
    il17 = oval(slide, Inches(5.90), Inches(3.30), Inches(1.50), Inches(0.95), CORE, CORE_DK)
    set_runs(il17, [(out_label, 13, True, WHITE)])

    mol_box = round_rect(slide, Inches(5.85), Inches(4.45), Inches(1.60), Inches(2.35), WHITE, MOL, adj=0.1)
    if lang == "fr":
        set_runs(mol_box, [
            (mol_title, 9, True, MOL),
            ("", 4, False, MUTED),
            ("podoplanine", 9, True, INK),
            ("(stroma ↔ immune)", 8, False, MUTED),
            ("", 4, False, MUTED),
            ("CD74 tronqué", 9, True, INK),
            ("(synoviocytes RA)", 8, False, MUTED),
            ("", 4, False, MUTED),
            ("ICAM-1 / VCAM-1", 9, True, INK),
            ("MMPs", 9, True, INK),
        ])
    else:
        set_runs(mol_box, [
            (mol_title, 9, True, MOL),
            ("", 4, False, MUTED),
            ("podoplanin", 9, True, INK),
            ("(stroma ↔ immune)", 8, False, MUTED),
            ("", 4, False, MUTED),
            ("truncated CD74", 9, True, INK),
            ("(RA synoviocytes)", 8, False, MUTED),
            ("", 4, False, MUTED),
            ("ICAM-1 / VCAM-1", 9, True, INK),
            ("MMPs", 9, True, INK),
        ])

    # Side arrows from left/right into center
    arrow_right(slide, Inches(5.40), Inches(2.00), Inches(0.28), Inches(0.20), POS)
    arrow_left(slide, Inches(7.62), Inches(2.00), Inches(0.28), Inches(0.20), NEG)

    # ===== RIGHT: negative interactions =====
    round_rect(slide, Inches(7.80), Inches(0.62), Inches(5.35), Inches(6.45), NEG_BAND, NEG, adj=0.03)
    neg_title = round_rect(slide, Inches(8.15), Inches(0.72), Inches(4.6), Inches(0.30), NEG, None, adj=0.2)
    set_runs(neg_title, [(neg_hdr, 11, True, WHITE)])

    # Treg ↔ Th17
    pair_n1 = round_rect(slide, Inches(7.97), Inches(1.12), Inches(5.00), Inches(1.85), WHITE, NEG, adj=0.06)
    textbox(slide, Inches(8.05), Inches(1.18), Inches(4.8), Inches(0.20), note_neg, size=9, bold=True, color=NEG_DK, align=PP_ALIGN.CENTER)
    cell(slide, Inches(8.20), Inches(1.50), Inches(1.30), Inches(1.15), TREG_LT, TREG, "Treg", treg_sub, 12, 8)
    contact_bridge(slide, Inches(9.60), Inches(1.75), Inches(1.20), Inches(0.28), "IL-10 / TGF-β")
    contact_bridge(slide, Inches(9.60), Inches(2.10), Inches(1.20), Inches(0.28), "CTLA-4 / PD-1")
    cell(slide, Inches(10.95), Inches(1.50), Inches(1.30), Inches(1.15), TH17_LT, TH17, "Th17", "↓ expansion", 12, 8)

    # M2 macrophages
    pair_n2 = round_rect(slide, Inches(7.97), Inches(3.10), Inches(5.00), Inches(1.55), WHITE, NEG, adj=0.06)
    textbox(
        slide, Inches(8.05), Inches(3.16), Inches(4.8), Inches(0.18),
        "Macrophages M2 (IL-4 / IL-13)" if lang == "fr" else "M2 macrophages (IL-4 / IL-13)",
        size=9, bold=True, color=NEG_DK, align=PP_ALIGN.CENTER,
    )
    cell(slide, Inches(8.30), Inches(3.45), Inches(1.35), Inches(1.00), M2_LT, M2, "M2", m2_sub, 12, 8)
    arrow_right(slide, Inches(9.80), Inches(3.80), Inches(0.40), Inches(0.22), NEG)
    cell(slide, Inches(10.40), Inches(3.45), Inches(1.35), Inches(1.00), TH17_LT, TH17, "Th17", "↓ IL-17", 12, 8)

    # Neutrophil feedback + pdpn-CLEC2
    pair_n3 = round_rect(slide, Inches(7.97), Inches(4.80), Inches(5.00), Inches(2.05), WHITE, NEG, adj=0.06)
    textbox(
        slide, Inches(8.05), Inches(4.86), Inches(4.8), Inches(0.18),
        "Feedback neutrophiles & pdpn–CLEC2" if lang == "fr" else "Neutrophil feedback & pdpn–CLEC2",
        size=9, bold=True, color=NEG_DK, align=PP_ALIGN.CENTER,
    )
    cell(slide, Inches(8.20), Inches(5.20), Inches(1.15), Inches(0.90), NEU_LT, NEU, "Neutro.", neu_sub, 10, 7)
    cell(slide, Inches(9.50), Inches(5.20), Inches(1.15), Inches(0.90), MAC_LT, MAC, "Macro.", "phagocytose" if lang == "fr" else "phagocytosis", 10, 7)
    arrow_right(slide, Inches(10.75), Inches(5.50), Inches(0.28), Inches(0.18), NEG)
    molecule_tag(slide, Inches(11.10), Inches(5.15), Inches(1.60), Inches(0.28), "↓ IL-23", NEG_LT, NEG)
    molecule_tag(slide, Inches(11.10), Inches(5.55), Inches(1.60), Inches(0.28), "↓ Th17", NEG_LT, NEG)
    molecule_tag(slide, Inches(8.20), Inches(6.25), Inches(4.50), Inches(0.40), clec_note, MOL_LT, MOL)

    textbox(
        slide, Inches(0.25), Inches(7.15), Inches(12.8), Inches(0.25),
        footer, size=8, color=MUTED, align=PP_ALIGN.CENTER,
    )


def build_legend_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    set_fill(bg, BG)
    no_line(bg)

    textbox(
        slide, Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.40),
        "Légende — Interactions cellulaires régulant l’IL-17",
        size=18, bold=True, color=INK, align=PP_ALIGN.LEFT,
    )
    legend = (
        "Cette figure illustre la régulation positive et négative de l’IL-17 par les interactions "
        "entre cellules immunitaires et stromales.\n\n"
        "À gauche (interactions activatrices), les cellules productrices d’IL-17 (Th17, γδ T, ILC3) "
        "entrent en contact avec les synoviocytes / fibroblastes. Les molécules de contact "
        "podoplanine et l’isoforme tronquée de CD74 favorisent la sécrétion d’IL-17. Les cellules "
        "dendritiques et macrophages fournissent IL-1β, IL-6 et IL-23, tandis qu’une boucle locale "
        "(IL-6, CXCL1/8, CCL20, ICAM-1/VCAM-1) entretient le recrutement et l’inflammation.\n\n"
        "À droite (interactions inhibitrices), les Treg freinent les Th17 via IL-10, TGF-β, CTLA-4 "
        "et PD-1 ; les macrophages M2 exercent un frein anti-inflammatoire. La phagocytose des "
        "neutrophiles apoptotiques diminue l’IL-23 et limite l’expansion Th17. L’interaction "
        "podoplanine–CLEC2 peut également réduire l’IL-17 et augmenter l’IL-10."
    )
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.95), Inches(12.0), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = legend
    run.font.size = Pt(14)
    run.font.color.rgb = INK
    run.font.name = "Calibri"


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    build_schema_slide(prs, lang="fr")
    build_schema_slide(prs, lang="en")
    build_legend_slide(prs)
    prs.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
