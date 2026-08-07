#!/usr/bin/env python3
"""Generate a Servier Medical Art–style PowerPoint figure for IL-17 regulation.

Layout:
  TOP    — Regulation of IL-17 production
  CENTER — IL-17 / IL-17RA–IL-17RC
  BOTTOM — Regulation of IL-17 function
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = "/workspace/figures/il17/Figure_IL17_production_fonction_Servier.pptx"

# Servier-like clean medical palette
BG = RGBColor(0xF8, 0xFA, 0xFC)
INK = RGBColor(0x2C, 0x3E, 0x50)
MUTED = RGBColor(0x5D, 0x6D, 0x7E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE_SOFT = RGBColor(0xCB, 0xD5, 0xE1)

PROD_HEADER = RGBColor(0x0E, 0x74, 0x86)
PROD_BAND = RGBColor(0xE6, 0xF4, 0xF6)
PROD_POS = RGBColor(0x0D, 0x94, 0x88)
PROD_POS_LT = RGBColor(0xE0, 0xF5, 0xF2)
PROD_CELL = RGBColor(0x3A, 0x7C, 0xA5)
PROD_CELL_LT = RGBColor(0xE4, 0xF0, 0xF8)
PROD_ENV = RGBColor(0x5B, 0x8F, 0x6B)
PROD_ENV_LT = RGBColor(0xE8, 0xF3, 0xEB)
PROD_NEG = RGBColor(0xD4, 0x6A, 0x4E)
PROD_NEG_LT = RGBColor(0xFB, 0xEE, 0xE9)

CORE = RGBColor(0xC4, 0x45, 0x69)
CORE_DK = RGBColor(0x8E, 0x2D, 0x4A)
CORE_LT = RGBColor(0xFB, 0xE9, 0xEF)
REC = RGBColor(0x7C, 0x3A, 0xED)
REC_LT = RGBColor(0xF1, 0xEA, 0xFE)
MEMBRANE = RGBColor(0x94, 0xA3, 0xB8)

FUNC_HEADER = RGBColor(0x2F, 0x4F, 0x78)
FUNC_BAND = RGBColor(0xEA, 0xEF, 0xF6)
FUNC_SYN = RGBColor(0xD9, 0x72, 0x16)
FUNC_SYN_LT = RGBColor(0xFE, 0xF3, 0xE6)
FUNC_BIO = RGBColor(0x4F, 0x5F, 0xB7)
FUNC_BIO_LT = RGBColor(0xEC, 0xEE, 0xFA)
FUNC_ANT = RGBColor(0x0F, 0x9A, 0x82)
FUNC_ANT_LT = RGBColor(0xE3, 0xF7, 0xF2)
FUNC_TISS = RGBColor(0x7A, 0x55, 0x48)
FUNC_TISS_LT = RGBColor(0xF4, 0xEC, 0xE9)


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def set_line(shape, color, width_pt=1.15):
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def round_rect(slide, left, top, width, height, fill, line=None, adj=0.12):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    set_fill(shape, fill)
    if line is None:
        no_line(shape)
    else:
        set_line(shape, line)
    try:
        shape.adjustments[0] = adj
    except Exception:
        pass
    return shape


def oval(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    set_fill(shape, fill)
    if line is None:
        no_line(shape)
    else:
        set_line(shape, line, 1.5)
    return shape


def down_arrow(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    set_fill(shape, color)
    no_line(shape)
    return shape


def set_runs(shape, lines, default_size=11, default_color=INK, align=PP_ALIGN.CENTER):
    """lines: list of (text, size, bold, color) or plain str."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    try:
        tf._txBody.get_or_add_bodyPr().set("anchor", "ctr")
    except Exception:
        pass
    for i, line in enumerate(lines):
        if isinstance(line, str):
            text, size, bold, color = line, default_size, False, default_color
        else:
            text, size, bold, color = line
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return tf


def textbox(slide, left, top, width, height, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def panel(slide, left, top, width, height, title, items, fill, accent, title_size=10, item_size=9):
    round_rect(slide, left, top, width, height, fill, accent, adj=0.1)
    bar = round_rect(slide, left, top, width, Inches(0.30), accent, None, adj=0.08)
    set_runs(bar, [(title, title_size, True, WHITE)])
    body = slide.shapes.add_textbox(
        left + Inches(0.10),
        top + Inches(0.34),
        width - Inches(0.18),
        height - Inches(0.40),
    )
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(1.5)
        run = p.add_run()
        run.text = "• " + item
        run.font.size = Pt(item_size)
        run.font.color.rgb = INK
        run.font.name = "Calibri"
    return body


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Background
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    set_fill(bg, BG)
    no_line(bg)

    # Title
    textbox(
        slide,
        Inches(0.3),
        Inches(0.08),
        Inches(12.7),
        Inches(0.34),
        "Régulation de la production et de la fonction de l’IL-17",
        size=20,
        bold=True,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    textbox(
        slide,
        Inches(0.3),
        Inches(0.38),
        Inches(12.7),
        Inches(0.24),
        "Schéma centré sur IL-17 / récepteurs — principes au-dessus (production), au-dessous (fonction)",
        size=11,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    # ---------- PRODUCTION BAND ----------
    round_rect(
        slide, Inches(0.22), Inches(0.68), Inches(12.9), Inches(2.78), PROD_BAND, PROD_HEADER, adj=0.03
    )
    hdr = round_rect(
        slide, Inches(4.15), Inches(0.76), Inches(5.0), Inches(0.32), PROD_HEADER, None, adj=0.22
    )
    set_runs(hdr, [("▲  RÉGULATION DE LA PRODUCTION", 12, True, WHITE)])

    pw, ph = Inches(2.95), Inches(2.10)
    py = Inches(1.18)
    gap = Inches(0.18)
    px0 = Inches(0.42)

    panel(
        slide,
        px0,
        py,
        pw,
        ph,
        "Cytokines inductrices (+)",
        [
            "TGF-β + IL-6 → RORγt / Th17",
            "IL-1β : phénotype inflammatoire",
            "IL-23 : maintien & pathogénicité",
            "IL-21 : induction indépendante",
        ],
        PROD_POS_LT,
        PROD_POS,
    )
    panel(
        slide,
        px0 + pw + gap,
        py,
        pw,
        ph,
        "Interactions cellulaires",
        [
            "Th17, γδ T, ILC3 ↔ stroma",
            "Synoviocytes / fibroblastes",
            "DC / macrophages (IL-1β/6/23)",
            "podoplanine & CD74 tronqué",
        ],
        PROD_CELL_LT,
        PROD_CELL,
    )
    panel(
        slide,
        px0 + 2 * (pw + gap),
        py,
        pw,
        ph,
        "Modulateurs environnementaux",
        [
            "Microbiote (SFB → Th17)",
            "Âge / inflammaging",
            "Nutrition & vitamine D",
            "Hormones (ex. œstrogènes)",
        ],
        PROD_ENV_LT,
        PROD_ENV,
    )
    panel(
        slide,
        px0 + 3 * (pw + gap),
        py,
        pw,
        ph,
        "Régulation négative (−)",
        [
            "IL-4 / IL-13 (STAT6–GATA3)",
            "IL-10 (Tregs, macrophages…)",
            "Balance Th17 / Treg",
            "Polarisation M2 anti-inflammatoire",
        ],
        PROD_NEG_LT,
        PROD_NEG,
    )

    for ax in (1.85, 5.0, 8.2, 11.35):
        down_arrow(slide, Inches(ax), Inches(3.48), Inches(0.26), Inches(0.28), PROD_HEADER)

    # ---------- CENTER: IL-17 / receptors ----------
    round_rect(
        slide, Inches(2.55), Inches(3.78), Inches(8.2), Inches(1.42), CORE_LT, CORE, adj=0.06
    )

    # Side context chips
    src = round_rect(
        slide, Inches(0.35), Inches(4.05), Inches(2.05), Inches(0.90), WHITE, LINE_SOFT, adj=0.12
    )
    set_runs(
        src,
        [
            ("Sources", 9, True, MUTED),
            ("Th17 • γδ T • ILC3", 10, True, INK),
        ],
    )
    tgt = round_rect(
        slide, Inches(10.95), Inches(4.05), Inches(2.05), Inches(0.90), WHITE, LINE_SOFT, adj=0.12
    )
    set_runs(
        tgt,
        [
            ("Cellules cibles", 9, True, MUTED),
            ("Stroma • épithélium", 10, True, INK),
            ("Endothélium • os…", 9, False, MUTED),
        ],
    )

    # Membrane-like bar (Servier topology cue)
    memb = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.35), Inches(4.95), Inches(3.25), Inches(0.14)
    )
    set_fill(memb, MEMBRANE)
    no_line(memb)
    try:
        memb.adjustments[0] = 0.5
    except Exception:
        pass

    il17 = oval(slide, Inches(2.85), Inches(3.98), Inches(2.55), Inches(1.05), CORE, CORE_DK)
    set_runs(
        il17,
        [
            ("IL-17A / IL-17F", 15, True, WHITE),
            ("(± hétérodimère A/F)", 9, False, RGBColor(0xFF, 0xD6, 0xE0)),
        ],
    )

    receptor = round_rect(
        slide, Inches(5.70), Inches(3.98), Inches(4.70), Inches(1.05), REC_LT, REC, adj=0.1
    )
    set_runs(
        receptor,
        [
            ("IL-17RA + IL-17RC", 14, True, REC),
            ("ACT1 → NF-κB / MAPK / C/EBP", 10, False, MUTED),
        ],
    )

    for ax in (1.85, 5.0, 8.2, 11.35):
        down_arrow(slide, Inches(ax), Inches(5.25), Inches(0.26), Inches(0.26), FUNC_HEADER)

    # ---------- FUNCTION BAND ----------
    round_rect(
        slide, Inches(0.22), Inches(5.50), Inches(12.9), Inches(1.72), FUNC_BAND, FUNC_HEADER, adj=0.03
    )
    fhdr = round_rect(
        slide, Inches(4.15), Inches(5.57), Inches(5.0), Inches(0.30), FUNC_HEADER, None, adj=0.22
    )
    set_runs(fhdr, [("▼  RÉGULATION DE LA FONCTION", 12, True, WHITE)])

    fy, fh = Inches(5.95), Inches(1.12)
    panel(
        slide,
        px0,
        fy,
        pw,
        fh,
        "Synergies cytokiniques",
        [
            "TNF via ↑ TNFRII (CUX1, IκBζ)",
            "IL-1β / IL-22 → MMPs, CCL20",
        ],
        FUNC_SYN_LT,
        FUNC_SYN,
        item_size=8.5,
    )
    panel(
        slide,
        px0 + pw + gap,
        fy,
        pw,
        fh,
        "Biodisponibilité",
        [
            "Auto-anticorps anti-IL-17 naturels",
            "Complexes Ag–Ac & biomarqueurs",
        ],
        FUNC_BIO_LT,
        FUNC_BIO,
        item_size=8.5,
    )
    panel(
        slide,
        px0 + 2 * (pw + gap),
        fy,
        pw,
        fh,
        "Antagonisme endogène",
        [
            "IL-25 (IL-17E) : IL-17RA / RB",
            "Compétition pour IL-17RA",
        ],
        FUNC_ANT_LT,
        FUNC_ANT,
        item_size=8.5,
    )
    panel(
        slide,
        px0 + 3 * (pw + gap),
        fy,
        pw,
        fh,
        "Contexte tissulaire",
        [
            "Intestin • vaisseaux • os",
            "Protection ↔ pathologie locale",
        ],
        FUNC_TISS_LT,
        FUNC_TISS,
        item_size=8.5,
    )

    textbox(
        slide,
        Inches(0.3),
        Inches(7.22),
        Inches(12.7),
        Inches(0.22),
        "D’après Noack & Miossec — Interactions cellulaires et moléculaires régulant la production et la fonction de l’IL-17  |  Topologie type illustration médicale Servier",
        size=8,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    # ---------- Slide 2: English twin for international use ----------
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    bg2 = slide2.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    set_fill(bg2, BG)
    no_line(bg2)

    textbox(
        slide2,
        Inches(0.3),
        Inches(0.08),
        Inches(12.7),
        Inches(0.34),
        "Regulation of IL-17 production and function",
        size=20,
        bold=True,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    textbox(
        slide2,
        Inches(0.3),
        Inches(0.38),
        Inches(12.7),
        Inches(0.24),
        "Centered on IL-17 / receptors — principles above (production), below (function)",
        size=11,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    round_rect(
        slide2, Inches(0.22), Inches(0.68), Inches(12.9), Inches(2.78), PROD_BAND, PROD_HEADER, adj=0.03
    )
    hdr2 = round_rect(
        slide2, Inches(4.15), Inches(0.76), Inches(5.0), Inches(0.32), PROD_HEADER, None, adj=0.22
    )
    set_runs(hdr2, [("▲  REGULATION OF PRODUCTION", 12, True, WHITE)])

    panel(
        slide2,
        px0,
        py,
        pw,
        ph,
        "Inducing cytokines (+)",
        [
            "TGF-β + IL-6 → RORγt / Th17",
            "IL-1β: inflammatory phenotype",
            "IL-23: maintenance & pathogenicity",
            "IL-21: IL-23-independent induction",
        ],
        PROD_POS_LT,
        PROD_POS,
    )
    panel(
        slide2,
        px0 + pw + gap,
        py,
        pw,
        ph,
        "Cellular interactions",
        [
            "Th17, γδ T, ILC3 ↔ stroma",
            "Synoviocytes / fibroblasts",
            "DC / macrophages (IL-1β/6/23)",
            "Podoplanin & truncated CD74",
        ],
        PROD_CELL_LT,
        PROD_CELL,
    )
    panel(
        slide2,
        px0 + 2 * (pw + gap),
        py,
        pw,
        ph,
        "Environmental modulators",
        [
            "Microbiota (SFB → Th17)",
            "Aging / inflammaging",
            "Nutrition & vitamin D",
            "Hormones (e.g. estrogens)",
        ],
        PROD_ENV_LT,
        PROD_ENV,
    )
    panel(
        slide2,
        px0 + 3 * (pw + gap),
        py,
        pw,
        ph,
        "Negative regulation (−)",
        [
            "IL-4 / IL-13 (STAT6–GATA3)",
            "IL-10 (Tregs, macrophages…)",
            "Th17 / Treg balance",
            "M2 macrophage polarization",
        ],
        PROD_NEG_LT,
        PROD_NEG,
    )

    for ax in (1.85, 5.0, 8.2, 11.35):
        down_arrow(slide2, Inches(ax), Inches(3.48), Inches(0.26), Inches(0.28), PROD_HEADER)

    round_rect(
        slide2, Inches(2.55), Inches(3.78), Inches(8.2), Inches(1.42), CORE_LT, CORE, adj=0.06
    )
    src2 = round_rect(
        slide2, Inches(0.35), Inches(4.05), Inches(2.05), Inches(0.90), WHITE, LINE_SOFT, adj=0.12
    )
    set_runs(src2, [("Sources", 9, True, MUTED), ("Th17 • γδ T • ILC3", 10, True, INK)])
    tgt2 = round_rect(
        slide2, Inches(10.95), Inches(4.05), Inches(2.05), Inches(0.90), WHITE, LINE_SOFT, adj=0.12
    )
    set_runs(
        tgt2,
        [
            ("Target cells", 9, True, MUTED),
            ("Stroma • epithelium", 10, True, INK),
            ("Endothelium • bone…", 9, False, MUTED),
        ],
    )
    memb2 = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.35), Inches(4.95), Inches(3.25), Inches(0.14)
    )
    set_fill(memb2, MEMBRANE)
    no_line(memb2)

    il172 = oval(slide2, Inches(2.85), Inches(3.98), Inches(2.55), Inches(1.05), CORE, CORE_DK)
    set_runs(
        il172,
        [
            ("IL-17A / IL-17F", 15, True, WHITE),
            ("(± A/F heterodimer)", 9, False, RGBColor(0xFF, 0xD6, 0xE0)),
        ],
    )
    receptor2 = round_rect(
        slide2, Inches(5.70), Inches(3.98), Inches(4.70), Inches(1.05), REC_LT, REC, adj=0.1
    )
    set_runs(
        receptor2,
        [
            ("IL-17RA + IL-17RC", 14, True, REC),
            ("ACT1 → NF-κB / MAPK / C/EBP", 10, False, MUTED),
        ],
    )

    for ax in (1.85, 5.0, 8.2, 11.35):
        down_arrow(slide2, Inches(ax), Inches(5.25), Inches(0.26), Inches(0.26), FUNC_HEADER)

    round_rect(
        slide2, Inches(0.22), Inches(5.50), Inches(12.9), Inches(1.72), FUNC_BAND, FUNC_HEADER, adj=0.03
    )
    fhdr2 = round_rect(
        slide2, Inches(4.15), Inches(5.57), Inches(5.0), Inches(0.30), FUNC_HEADER, None, adj=0.22
    )
    set_runs(fhdr2, [("▼  REGULATION OF FUNCTION", 12, True, WHITE)])

    panel(
        slide2,
        px0,
        fy,
        pw,
        fh,
        "Cytokine synergies",
        [
            "TNF via ↑ TNFRII (CUX1, IκBζ)",
            "IL-1β / IL-22 → MMPs, CCL20",
        ],
        FUNC_SYN_LT,
        FUNC_SYN,
        item_size=8.5,
    )
    panel(
        slide2,
        px0 + pw + gap,
        fy,
        pw,
        fh,
        "Bioavailability",
        [
            "Natural anti–IL-17 autoantibodies",
            "Ag–Ab complexes & biomarkers",
        ],
        FUNC_BIO_LT,
        FUNC_BIO,
        item_size=8.5,
    )
    panel(
        slide2,
        px0 + 2 * (pw + gap),
        fy,
        pw,
        fh,
        "Endogenous antagonism",
        [
            "IL-25 (IL-17E): IL-17RA / RB",
            "Competition for IL-17RA",
        ],
        FUNC_ANT_LT,
        FUNC_ANT,
        item_size=8.5,
    )
    panel(
        slide2,
        px0 + 3 * (pw + gap),
        fy,
        pw,
        fh,
        "Tissue context",
        [
            "Intestine • vessels • bone",
            "Protection ↔ local pathology",
        ],
        FUNC_TISS_LT,
        FUNC_TISS,
        item_size=8.5,
    )

    textbox(
        slide2,
        Inches(0.3),
        Inches(7.22),
        Inches(12.7),
        Inches(0.22),
        "Based on Noack & Miossec — Cellular and molecular interactions regulating IL-17 production and function  |  Servier-style medical topology",
        size=8,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    # ---------- Slide 3: short figure legend (FR) ----------
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    bg3 = slide3.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    set_fill(bg3, BG)
    no_line(bg3)
    textbox(
        slide3,
        Inches(0.6),
        Inches(0.4),
        Inches(12.0),
        Inches(0.4),
        "Légende — Figure IL-17 (production / récepteurs / fonction)",
        size=18,
        bold=True,
        color=INK,
        align=PP_ALIGN.LEFT,
    )
    legend = (
        "Ce schéma distingue deux niveaux de régulation de l’IL-17. "
        "Au-dessus, la régulation de la production intègre les cytokines inductrices "
        "(TGF-β, IL-6, IL-1β, IL-23, IL-21), les interactions cellulaires "
        "(Th17/γδ T/ILC3 avec le stroma ; podoplanine et isoforme tronquée de CD74), "
        "les facteurs environnementaux (microbiote, âge, nutrition, hormones) "
        "et les freins négatifs (IL-4/IL-13, IL-10, balance Th17/Treg, M2).\n\n"
        "Au centre, IL-17A/F engage le complexe IL-17RA/IL-17RC et la cascade "
        "ACT1–NF-κB/MAPK/C/EBP.\n\n"
        "Au-dessous, la régulation de la fonction porte sur les synergies "
        "(notamment TNF via TNFRII), la biodisponibilité (auto-anticorps anti-IL-17), "
        "l’antagonisme endogène (IL-25/IL-17E) et le contexte tissulaire "
        "(intestin, vaisseaux, os), où l’IL-17 peut être protectrice ou pathologique."
    )
    box = slide3.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12.0), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = legend
    run.font.size = Pt(14)
    run.font.color.rgb = INK
    run.font.name = "Calibri"

    prs.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
