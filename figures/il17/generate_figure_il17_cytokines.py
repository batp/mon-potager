#!/usr/bin/env python3
"""Servier-style PPTX: positive vs negative cytokine regulation of IL-17.

Based on Noack & Miossec — section on cytokine interactions regulating
IL-17 production and function.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

OUT = "/workspace/figures/il17/Figure_IL17_cytokines_pos_neg.pptx"

BG = RGBColor(0xF8, 0xFA, 0xFC)
INK = RGBColor(0x2C, 0x3E, 0x50)
MUTED = RGBColor(0x5D, 0x6D, 0x7E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xCB, 0xD5, 0xE1)

POS = RGBColor(0x0D, 0x94, 0x88)
POS_DK = RGBColor(0x0A, 0x6F, 0x66)
POS_LT = RGBColor(0xE0, 0xF5, 0xF2)
POS_BAND = RGBColor(0xE8, 0xF7, 0xF4)

NEG = RGBColor(0xD4, 0x6A, 0x4E)
NEG_DK = RGBColor(0xA8, 0x4A, 0x32)
NEG_LT = RGBColor(0xFB, 0xEE, 0xE9)
NEG_BAND = RGBColor(0xFC, 0xF1, 0xEC)

CORE = RGBColor(0xC4, 0x45, 0x69)
CORE_DK = RGBColor(0x8E, 0x2D, 0x4A)
CORE_LT = RGBColor(0xFB, 0xE9, 0xEF)
REC = RGBColor(0x6D, 0x28, 0xD9)
REC_LT = RGBColor(0xF1, 0xEA, 0xFE)
CHIP = RGBColor(0xEE, 0xF2, 0xF7)


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def set_line(shape, color, width_pt=1.15):
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def round_rect(slide, left, top, width, height, fill, line=None, adj=0.12):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    set_fill(shape, fill)
    if line is None:
        no_line(shape)
    else:
        set_line(shape, line)
    try:
        shape.adjustments[0] = adj
    except Exception:
        pass
    return shape


def oval(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    set_fill(shape, fill)
    if line is None:
        no_line(shape)
    else:
        set_line(shape, line, 1.5)
    return shape


def chevron_right(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    set_fill(shape, color)
    no_line(shape)
    return shape


def chevron_left(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, left, top, width, height)
    set_fill(shape, color)
    no_line(shape)
    return shape


def set_runs(shape, lines, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    try:
        tf._txBody.get_or_add_bodyPr().set("anchor", "ctr")
    except Exception:
        pass
    for i, line in enumerate(lines):
        if isinstance(line, str):
            text, size, bold, color = line, 11, False, INK
        else:
            text, size, bold, color = line
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return tf


def textbox(slide, left, top, width, height, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def cytokine_chip(slide, left, top, width, height, name, subtitle, fill, accent):
    shape = round_rect(slide, left, top, width, height, fill, accent, adj=0.18)
    set_runs(
        shape,
        [
            (name, 12, True, accent),
            (subtitle, 8.5, False, MUTED),
        ],
    )
    return shape


def section_card(slide, left, top, width, height, title, bullets, fill, accent, item_size=9):
    round_rect(slide, left, top, width, height, fill, accent, adj=0.08)
    bar = round_rect(slide, left, top, width, Inches(0.32), accent, None, adj=0.08)
    set_runs(bar, [(title, 11, True, WHITE)])
    body = slide.shapes.add_textbox(
        left + Inches(0.12),
        top + Inches(0.38),
        width - Inches(0.22),
        height - Inches(0.46),
    )
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(3)
        run = p.add_run()
        run.text = "• " + item
        run.font.size = Pt(item_size)
        run.font.color.rgb = INK
        run.font.name = "Calibri"
    return body


def build_schema_slide(prs, lang="fr"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    set_fill(bg, BG)
    no_line(bg)

    if lang == "fr":
        title = "Interactions cytokiniques régulant l’IL-17"
        subtitle = "Régulation positive (gauche) et négative (droite) — production et fonction"
        pos_hdr = "+  RÉGULATION POSITIVE"
        neg_hdr = "−  RÉGULATION NÉGATIVE"
        prod_pos_title = "Production ↑"
        prod_neg_title = "Production ↓"
        func_pos_title = "Fonction ↑"
        func_neg_title = "Fonction ↓"
        prod_pos = [
            "TGF-β + IL-6 → RORγt / Th17",
            "IL-1β renforce le phénotype inflammatoire",
            "IL-23 : maintien, survie & pathogénicité",
            "IL-21 : induction indépendante d’IL-23",
        ]
        func_pos = [
            "Synergie IL-17 + TNF (↑ TNFRII)",
            "CUX1 & IκBζ avec NF-κB",
            "Coopération avec IL-1β et IL-22",
            "↑ chimiokines, MMPs, CCL20",
        ]
        prod_neg = [
            "IL-4 : STAT6–GATA3 antagonise RORγt",
            "IL-13 (axe Th2) freine Th17",
            "IL-10 ↓ IL-1β / IL-6 / IL-23 des APC",
            "IL-10 agit aussi sur les effecteurs T",
        ]
        func_neg = [
            "IL-25 (IL-17E) : IL-17RA + IL-17RB",
            "Compétition pour la chaîne IL-17RA",
            "Auto-Ac naturels anti-IL-17",
            "Neutralisation de l’IL-17 biodisponible",
        ]
        center_src = "Naïve CD4⁺ → Th17"
        center_out = "IL-17A / IL-17F"
        center_rec = "IL-17RA + IL-17RC"
        center_sig = "ACT1 → NF-κB / MAPK"
        footer = (
            "D’après Noack & Miossec — Interactions avec les cytokines "
            "(régulation positive / négative de la production et de la fonction)  |  Style Servier"
        )
        key_pos = "Inducteurs & amplificateurs"
        key_neg = "Inhibiteurs & antagonistes"
    else:
        title = "Cytokine interactions regulating IL-17"
        subtitle = "Positive (left) and negative (right) regulation — production and function"
        pos_hdr = "+  POSITIVE REGULATION"
        neg_hdr = "−  NEGATIVE REGULATION"
        prod_pos_title = "Production ↑"
        prod_neg_title = "Production ↓"
        func_pos_title = "Function ↑"
        func_neg_title = "Function ↓"
        prod_pos = [
            "TGF-β + IL-6 → RORγt / Th17",
            "IL-1β enhances inflammatory phenotype",
            "IL-23: maintenance, survival & pathogenicity",
            "IL-21: IL-23-independent induction",
        ]
        func_pos = [
            "IL-17 + TNF synergy (↑ TNFRII)",
            "CUX1 & IκBζ with NF-κB",
            "Cooperation with IL-1β and IL-22",
            "↑ chemokines, MMPs, CCL20",
        ]
        prod_neg = [
            "IL-4: STAT6–GATA3 antagonizes RORγt",
            "IL-13 (Th2 axis) restrains Th17",
            "IL-10 ↓ APC IL-1β / IL-6 / IL-23",
            "IL-10 also acts on effector T cells",
        ]
        func_neg = [
            "IL-25 (IL-17E): IL-17RA + IL-17RB",
            "Competition for shared IL-17RA",
            "Natural anti–IL-17 autoantibodies",
            "Neutralization of bioactive IL-17",
        ]
        center_src = "Naïve CD4⁺ → Th17"
        center_out = "IL-17A / IL-17F"
        center_rec = "IL-17RA + IL-17RC"
        center_sig = "ACT1 → NF-κB / MAPK"
        footer = (
            "Based on Noack & Miossec — Cytokine interactions "
            "(positive / negative regulation of production and function)  |  Servier style"
        )
        key_pos = "Inducers & amplifiers"
        key_neg = "Inhibitors & antagonists"

    textbox(
        slide, Inches(0.3), Inches(0.08), Inches(12.7), Inches(0.34),
        title, size=20, bold=True, color=INK, align=PP_ALIGN.CENTER,
    )
    textbox(
        slide, Inches(0.3), Inches(0.38), Inches(12.7), Inches(0.24),
        subtitle, size=11, color=MUTED, align=PP_ALIGN.CENTER,
    )

    # Left positive band
    round_rect(slide, Inches(0.22), Inches(0.70), Inches(4.55), Inches(6.35), POS_BAND, POS, adj=0.03)
    pos_title = round_rect(slide, Inches(0.55), Inches(0.82), Inches(3.9), Inches(0.34), POS, None, adj=0.2)
    set_runs(pos_title, [(pos_hdr, 12, True, WHITE)])
    textbox(slide, Inches(0.45), Inches(1.20), Inches(4.1), Inches(0.22), key_pos, size=9, color=POS_DK, align=PP_ALIGN.CENTER)

    # Positive cytokine chips (top row)
    chips = [
        ("IL-6", "avec TGF-β" if lang == "fr" else "with TGF-β"),
        ("IL-1β", "différenciation" if lang == "fr" else "differentiation"),
        ("IL-23", "pathogénicité" if lang == "fr" else "pathogenicity"),
        ("IL-21", "induction"),
        ("TNF", "synergie" if lang == "fr" else "synergy"),
        ("IL-22", "amplification"),
    ]
    chip_w, chip_h = Inches(1.35), Inches(0.55)
    for i, (name, sub) in enumerate(chips):
        col, row = i % 3, i // 3
        cytokine_chip(
            slide,
            Inches(0.45) + col * Inches(1.45),
            Inches(1.48) + row * Inches(0.62),
            chip_w,
            chip_h,
            name,
            sub,
            POS_LT,
            POS,
        )

    section_card(
        slide, Inches(0.42), Inches(2.85), Inches(4.15), Inches(1.95),
        prod_pos_title, prod_pos, WHITE, POS, item_size=9.5,
    )
    section_card(
        slide, Inches(0.42), Inches(4.95), Inches(4.15), Inches(1.90),
        func_pos_title, func_pos, WHITE, POS_DK, item_size=9.5,
    )

    # Right negative band
    round_rect(slide, Inches(8.55), Inches(0.70), Inches(4.55), Inches(6.35), NEG_BAND, NEG, adj=0.03)
    neg_title = round_rect(slide, Inches(8.88), Inches(0.82), Inches(3.9), Inches(0.34), NEG, None, adj=0.2)
    set_runs(neg_title, [(neg_hdr, 12, True, WHITE)])
    textbox(slide, Inches(8.78), Inches(1.20), Inches(4.1), Inches(0.22), key_neg, size=9, color=NEG_DK, align=PP_ALIGN.CENTER)

    neg_chips = [
        ("IL-4", "STAT6–GATA3"),
        ("IL-13", "Th2"),
        ("IL-10", "Treg / APC"),
        ("IL-25", "IL-17E"),
        ("anti–IL-17", "auto-Ac" if lang == "fr" else "auto-Ab"),
        ("TGF-β*", "Treg si −IL-6" if lang == "fr" else "Treg if −IL-6"),
    ]
    for i, (name, sub) in enumerate(neg_chips):
        col, row = i % 3, i // 3
        cytokine_chip(
            slide,
            Inches(8.78) + col * Inches(1.45),
            Inches(1.48) + row * Inches(0.62),
            chip_w,
            chip_h,
            name,
            sub,
            NEG_LT,
            NEG,
        )

    section_card(
        slide, Inches(8.75), Inches(2.85), Inches(4.15), Inches(1.95),
        prod_neg_title, prod_neg, WHITE, NEG, item_size=9.5,
    )
    section_card(
        slide, Inches(8.75), Inches(4.95), Inches(4.15), Inches(1.90),
        func_neg_title, func_neg, WHITE, NEG_DK, item_size=9.5,
    )

    # Center axis
    round_rect(slide, Inches(5.00), Inches(0.70), Inches(3.30), Inches(6.35), CORE_LT, CORE, adj=0.04)

    hub_title = round_rect(slide, Inches(5.25), Inches(0.90), Inches(2.80), Inches(0.34), CORE, None, adj=0.2)
    set_runs(hub_title, [("IL-17", 13, True, WHITE)])

    src = round_rect(slide, Inches(5.25), Inches(1.40), Inches(2.80), Inches(0.70), WHITE, LINE, adj=0.12)
    set_runs(src, [(center_src, 11, True, INK), ("RORγt", 10, False, MUTED)])

    # Positive arrows into center
    chevron_right(slide, Inches(4.78), Inches(3.35), Inches(0.28), Inches(0.22), POS)
    chevron_right(slide, Inches(4.78), Inches(5.55), Inches(0.28), Inches(0.22), POS)
    # Negative arrows into center
    chevron_left(slide, Inches(8.25), Inches(3.35), Inches(0.28), Inches(0.22), NEG)
    chevron_left(slide, Inches(8.25), Inches(5.55), Inches(0.28), Inches(0.22), NEG)

    il17 = oval(slide, Inches(5.35), Inches(2.35), Inches(2.60), Inches(1.15), CORE, CORE_DK)
    set_runs(il17, [(center_out, 14, True, WHITE), ("cytokine effectrice", 9, False, RGBColor(0xFF, 0xD6, 0xE0)) if lang == "fr" else ("effector cytokine", 9, False, RGBColor(0xFF, 0xD6, 0xE0))])

    rec = round_rect(slide, Inches(5.25), Inches(3.70), Inches(2.80), Inches(1.05), REC_LT, REC, adj=0.12)
    set_runs(rec, [(center_rec, 12, True, REC), (center_sig, 9, False, MUTED)])

    note = round_rect(slide, Inches(5.25), Inches(4.95), Inches(2.80), Inches(1.80), WHITE, LINE, adj=0.1)
    if lang == "fr":
        note_lines = [
            ("Principe", 10, True, CORE),
            ("Les cytokines contrôlent", 9, False, INK),
            ("à la fois la production", 9, False, INK),
            ("(axe Th17) et la fonction", 9, False, INK),
            ("(synergie / antagonisme", 9, False, INK),
            ("au niveau du récepteur).", 9, False, INK),
            ("", 6, False, MUTED),
            ("* TGF-β seul → Treg", 8, False, MUTED),
        ]
    else:
        note_lines = [
            ("Principle", 10, True, CORE),
            ("Cytokines control both", 9, False, INK),
            ("production (Th17 axis)", 9, False, INK),
            ("and function (receptor-", 9, False, INK),
            ("level synergy /", 9, False, INK),
            ("antagonism).", 9, False, INK),
            ("", 6, False, MUTED),
            ("* TGF-β alone → Treg", 8, False, MUTED),
        ]
    set_runs(note, note_lines)

    textbox(
        slide, Inches(0.3), Inches(7.15), Inches(12.7), Inches(0.25),
        footer, size=8, color=MUTED, align=PP_ALIGN.CENTER,
    )


def build_legend_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    set_fill(bg, BG)
    no_line(bg)

    textbox(
        slide, Inches(0.6), Inches(0.4), Inches(12.0), Inches(0.4),
        "Légende — Interactions cytokiniques et régulation de l’IL-17",
        size=18, bold=True, color=INK, align=PP_ALIGN.LEFT,
    )
    legend = (
        "Cette figure résume la régulation positive et négative de l’IL-17 par les interactions cytokiniques, "
        "en distinguant production et fonction.\n\n"
        "À gauche (régulation positive), TGF-β et IL-6 initient la différenciation Th17 via RORγt ; "
        "IL-1β renforce le phénotype inflammatoire ; IL-23 assure le maintien et la pathogénicité ; "
        "IL-21 peut induire Th17 indépendamment d’IL-23. Au niveau fonctionnel, la synergie IL-17/TNF "
        "(induction de TNFRII, CUX1, IκBζ) et la coopération avec IL-1β et IL-22 amplifient chimiokines, "
        "MMPs et CCL20.\n\n"
        "À droite (régulation négative), IL-4 et IL-13 freinent Th17 via STAT6–GATA3 ; IL-10 diminue "
        "IL-1β/IL-6/IL-23 des APC et limite les effecteurs T. IL-25 (IL-17E) antagonise le signal "
        "IL-17A/F en compétition pour IL-17RA, tandis que les auto-anticorps naturels anti-IL-17 "
        "réduisent la biodisponibilité de la cytokine."
    )
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12.0), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = legend
    run.font.size = Pt(14)
    run.font.color.rgb = INK
    run.font.name = "Calibri"


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    build_schema_slide(prs, lang="fr")
    build_schema_slide(prs, lang="en")
    build_legend_slide(prs)
    prs.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
